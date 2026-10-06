import io
import os
import re
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import HTMLResponse, StreamingResponse
from PIL import Image
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from google import genai
from google.genai import types
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    HRFlowable
)
from reportlab.lib.units import cm

app = FastAPI(title="Cirugía Léxica ELE")

API_KEY = os.environ.get("GEMINI_API_KEY")
if not API_KEY:
    print("ADVERTENCIA: La variable de entorno GEMINI_API_KEY no está configurada.")

cliente_gemini = genai.Client(api_key=API_KEY)


def limpiar_markdown_residual(texto: str) -> str:
    linea = texto.strip()
    if linea.startswith("```"):
        return ""
    return linea


def fabricar_word(texto_ia: str, contexto_destino: str) -> io.BytesIO:
    doc = Document()

    for seccion in doc.sections:
        seccion.top_margin = Inches(0.8)
        seccion.bottom_margin = Inches(0.8)
        seccion.left_margin = Inches(0.9)
        seccion.right_margin = Inches(0.9)

    tag = "ACTIVIDAD DIDÁCTICA ADAPTADA • CONTEXTO: " + contexto_destino.upper()
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(4)
    run_meta = p_meta.add_run(tag)
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(8.5)
    run_meta.font.bold = True
    run_meta.font.color.rgb = RGBColor(120, 120, 120)

    p_alumno = doc.add_paragraph()
    p_alumno.paragraph_format.space_after = Pt(16)
    subraya = "_" * 30
    run_alumno = p_alumno.add_run(f"Nombre del estudiante: {subraya}   Fecha: ___________")
    run_alumno.font.name = "Arial"
    run_alumno.font.size = Pt(9.5)
    run_alumno.font.color.rgb = RGBColor(90, 90, 90)

    texto_reparado = re.sub(r"\n\s*\|\s*", " | ", texto_ia)
    lineas = [l.strip() for l in texto_reparado.split("\n") if l.strip()]
    
    i = 0
    total_lineas = len(lineas)

    while i < total_lineas:
        linea_raw = lineas[i]
        linea_str = limpiar_markdown_residual(linea_raw)

        if not linea_str:
            i += 1
            continue

        if "|" in linea_raw and not linea_raw.startswith(("#", "Respuestas:", "Respuesta:")):
            filas_tabla = []
            while i < total_lineas and "|" in lineas[i] and not lineas[i].startswith(("#", "Respuestas:", "Respuesta:")):
                fila_actual = lineas[i]
                if not re.match(r"^\|?(\s*:?-+:?\s*\|?)+$", fila_actual):
                    celdas = [c.strip() for c in fila_actual.strip("|").split("|") if c.strip()]
                    if len(celdas) >= 2:
                        filas_tabla.append(celdas[:2])
                i += 1

            if filas_tabla:
                tabla = doc.add_table(rows=len(filas_tabla), cols=2)
                tabla.style = "Table Grid"

                for num_f, fila in enumerate(filas_tabla):
                    for num_c, texto_celda in enumerate(fila):
                        celda = tabla.cell(num_f, num_c)
                        celda.text = ""
                        p_celda = celda.paragraphs[0]
                        p_celda.paragraph_format.space_before = Pt(3)
                        p_celda.paragraph_format.space_after = Pt(3)

                        partes = re.split(r"(\*\*.*?\*\*)", texto_celda)
                        for parte in partes:
                            if parte.startswith("**") and parte.endswith("**"):
                                r = p_celda.add_run(parte[2:-2])
                                r.font.bold = True
                            else:
                                r = p_celda.add_run(re.sub(r"[*_]", "", parte))
                            r.font.name = "Arial"
                            r.font.size = Pt(9.5)
                            if num_f == 0:
                                r.font.bold = True

                p_espacio = doc.add_paragraph()
                p_espacio.paragraph_format.space_before = Pt(2)
                p_espacio.paragraph_format.space_after = Pt(4)
            continue

        if linea_str in ("---", "***", "___"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        if linea_str.lower().startswith(("foto:", "pie de foto", "imagen:")):
            i += 1
            continue

        if re.match(r"^#\s+", linea_raw):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            texto_titulo = re.sub(r"^#+\s*", "", linea_str).strip()
            run = p.add_run(re.sub(r"[*_]", "", texto_titulo))
            run.font.name = "Arial"
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(24, 43, 73)

        elif re.match(r"^##\s+", linea_raw):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            texto_subtitulo = re.sub(r"^#+\s*", "", linea_str).strip()
            run = p.add_run(re.sub(r"[*_]", "", texto_subtitulo))
            run.font.name = "Arial"
            run.font.size = Pt(12)
            run.font.bold